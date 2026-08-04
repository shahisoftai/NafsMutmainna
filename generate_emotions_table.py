import json
import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))

def load_json(path):
    with open(os.path.join(BASE, path), 'r', encoding='utf-8') as f:
        return json.load(f)

# Load all data
emotions_data = load_json("assets/data/emotions_seed.json")["rows"]
quran_ayat = load_json("assets/data/quran_ayat_seed.json")["rows"]
hadees_data = load_json("assets/data/hadees_seed.json")["rows"]
duas_data = load_json("assets/data/duas_seed.json")["duas"]
eq_links = load_json("assets/data/emotion_quran_links_seed.json")["rows"]
eh_links = load_json("assets/data/emotion_hadees_links_seed.json")["rows"]

# Build lookup maps
quran_map = {a["Ayat_ID"]: a for a in quran_ayat}
hadees_map = {h["Hadees_ID"]: h for h in hadees_data}
duas_by_emotion = {}
for d in duas_data:
    eid = d["Emotion_ID"]
    duas_by_emotion.setdefault(eid, []).append(d)

quran_by_emotion = {}
for link in eq_links:
    eid = link["Emotion_ID"]
    quran_by_emotion.setdefault(eid, []).append(link)

hadees_by_emotion = {}
for link in eh_links:
    eid = link["Emotion_ID"]
    hadees_by_emotion.setdefault(eid, []).append(link)

def set_cell_shading(cell, color):
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color
    })
    shading.append(shading_elem)

def add_cell_text(cell, text, bold=False, size=9, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

doc = Document()

# Title
title = doc.add_heading("HeartOS — Database of 50 Core Emotions", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Complete Reference Table: Emotions with Related Quran, Hadees, Duas, Allah Names & More")
run.font.size = Pt(11)
run.italic = True

doc.add_paragraph()

# Count emotions
neg_count = sum(1 for e in emotions_data if e.get("Category") == "Negative")
pos_count = sum(1 for e in emotions_data if e.get("Category") == "Positive")
stats = doc.add_paragraph()
stats.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = stats.add_run(f"Total: {len(emotions_data)} emotions ({neg_count} Negative, {pos_count} Positive)")
run.font.size = Pt(10)

doc.add_paragraph()

# Build table for each emotion
for idx, em in enumerate(emotions_data):
    eid = em["Emotion_ID"]
    cat = em.get("Category", "")
    is_positive = cat == "Positive"
    
    # Section header
    h = doc.add_heading(f"{idx+1}. {em['Core_Emotion']} ({em.get('Arabic_Name','')})", level=1)
    
    # Info table
    info_table = doc.add_table(rows=5, cols=2)
    info_table.style = 'Table Grid'
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    info_data = [
        ("English", em["Core_Emotion"]),
        ("Arabic", em.get("Arabic_Name", "")),
        ("Urdu", em.get("Urdu_Name", "")),
        ("Category", cat),
        ("Description", em.get("Description", "")),
    ]
    
    for i, (label, value) in enumerate(info_data):
        cell0 = info_table.cell(i, 0)
        cell1 = info_table.cell(i, 1)
        add_cell_text(cell0, label, bold=True, size=9)
        set_cell_shading(cell0, "E8E8E8")
        add_cell_text(cell1, value, size=9)
        cell0.width = Inches(1.2)
        cell1.width = Inches(5.3)
    
    doc.add_paragraph()
    
    # Spiritual Interventions Table
    inter_heading = doc.add_heading("Spiritual Interventions", level=2)
    
    inter_table = doc.add_table(rows=5, cols=2)
    inter_table.style = 'Table Grid'
    inter_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    interventions = [
        ("Recommended Dua", em.get("Recommended_Dua_English", "")),
        ("Recommended Allah Names", em.get("Recommended_Allah_Names", "")),
        ("Recommended Dhikr", em.get("Recommended_Dhikr", "")),
        ("Daily Action", em.get("Daily_Action", "")),
        ("Growth Path", em.get("Growth_Path", "")),
    ]
    
    for i, (label, value) in enumerate(interventions):
        cell0 = inter_table.cell(i, 0)
        cell1 = inter_table.cell(i, 1)
        add_cell_text(cell0, label, bold=True, size=9)
        set_cell_shading(cell0, "E8F0FE")
        add_cell_text(cell1, value, size=9)
        cell0.width = Inches(1.8)
        cell1.width = Inches(4.7)
    
    # Dua Arabic if available
    if em.get("Recommended_Dua_Arabic"):
        doc.add_paragraph()
        dua_p = doc.add_paragraph()
        run = dua_p.add_run(f"Dua (Arabic): {em['Recommended_Dua_Arabic']}")
        run.font.size = Pt(9)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        ref = em.get("Recommended_Dua_Reference", "")
        if ref:
            run2 = dua_p.add_run(f"\nSource: {ref}")
            run2.font.size = Pt(8)
            run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    
    doc.add_paragraph()
    
    # Quran Verses
    quran_heading = doc.add_heading("Related Quran Verses", level=2)
    q_links = quran_by_emotion.get(eid, [])
    if q_links:
        q_table = doc.add_table(rows=1, cols=3)
        q_table.style = 'Table Grid'
        q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Header
        for ci, hdr in enumerate(["Reference", "Arabic", "English Translation"]):
            cell = q_table.cell(0, ci)
            add_cell_text(cell, hdr, bold=True, size=8)
            set_cell_shading(cell, "1B5E20")
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        for link in sorted(q_links, key=lambda x: x.get("Weight", 0), reverse=True):
            ayat = quran_map.get(link["Ayat_ID"])
            if not ayat:
                continue
            row = q_table.add_row()
            add_cell_text(row.cells[0], ayat.get("Full_Reference", ""), size=8)
            add_cell_text(row.cells[1], ayat.get("Arabic_Text", ""), size=8)
            add_cell_text(row.cells[2], ayat.get("English_Translation", ""), size=8)
        
        # Set column widths
        for row in q_table.rows:
            row.cells[0].width = Inches(1.3)
            row.cells[1].width = Inches(2.5)
            row.cells[2].width = Inches(2.7)
    else:
        doc.add_paragraph("(No specific Quran verses linked)").paragraph_format.space_before = Pt(2)
    
    doc.add_paragraph()
    
    # Hadees
    hadees_heading = doc.add_heading("Related Hadees", level=2)
    h_links = hadees_by_emotion.get(eid, [])
    if h_links:
        h_table = doc.add_table(rows=1, cols=3)
        h_table.style = 'Table Grid'
        h_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for ci, hdr in enumerate(["Source", "Arabic", "English Translation"]):
            cell = h_table.cell(0, ci)
            add_cell_text(cell, hdr, bold=True, size=8)
            set_cell_shading(cell, "1A237E")
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        for link in sorted(h_links, key=lambda x: x.get("Weight", 0), reverse=True):
            h = hadees_map.get(link["Hadees_ID"])
            if not h:
                continue
            row = h_table.add_row()
            source = f"{h.get('Source_Book', '')} {h.get('Hadith_Number', '')} ({h.get('Grade', '')})"
            add_cell_text(row.cells[0], source, size=8)
            add_cell_text(row.cells[1], h.get("Arabic_Text", ""), size=8)
            add_cell_text(row.cells[2], h.get("English_Translation", ""), size=8)
        
        for row in h_table.rows:
            row.cells[0].width = Inches(1.5)
            row.cells[1].width = Inches(2.5)
            row.cells[2].width = Inches(2.5)
    else:
        doc.add_paragraph("(No specific Hadees linked)").paragraph_format.space_before = Pt(2)
    
    doc.add_paragraph()
    
    # Additional Duas
    duas_heading = doc.add_heading("Additional Duas", level=2)
    extra_duas = duas_by_emotion.get(eid, [])
    if extra_duas:
        d_table = doc.add_table(rows=1, cols=3)
        d_table.style = 'Table Grid'
        d_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for ci, hdr in enumerate(["Source", "Arabic", "English"]):
            cell = d_table.cell(0, ci)
            add_cell_text(cell, hdr, bold=True, size=8)
            set_cell_shading(cell, "4A148C")
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        for d in extra_duas:
            row = d_table.add_row()
            ref = f"{d.get('Reference', '')} ({d.get('Source_Type', '')})"
            add_cell_text(row.cells[0], ref, size=8)
            add_cell_text(row.cells[1], d.get("Arabic", ""), size=8)
            add_cell_text(row.cells[2], d.get("English", ""), size=8)
        
        for row in d_table.rows:
            row.cells[0].width = Inches(1.5)
            row.cells[1].width = Inches(2.5)
            row.cells[2].width = Inches(2.5)
    else:
        doc.add_paragraph("(No additional duas)").paragraph_format.space_before = Pt(2)
    
    # Page break between emotions
    doc.add_page_break()

# Save
output_path = os.path.join(BASE, "HeartOS_50_Emotions_Complete_Reference.docx")
doc.save(output_path)
print(f"Document saved to: {output_path}")
print(f"Total emotions: {len(emotions_data)}")
print(f"Total pages estimated: ~{len(emotions_data) * 2}+")
