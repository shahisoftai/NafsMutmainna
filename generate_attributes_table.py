import json
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE = os.path.dirname(os.path.abspath(__file__))

def load_json(path):
    with open(os.path.join(BASE, path), 'r', encoding='utf-8') as f:
        return json.load(f)

emotions = load_json("assets/data/emotions_seed.json")["rows"]
attrs = load_json("assets/data/attributes_seed.json")["rows"]
ea_links = load_json("assets/data/emotion_attribute_links_seed.json")["rows"]

attr_map = {a["Attribute_ID"]: a for a in attrs}

links_by_emotion = {}
for link in ea_links:
    eid = link["Emotion_ID"]
    links_by_emotion.setdefault(eid, []).append(link)

ROLE_ORDER = {"Core": 0, "Disease": 1, "Treatment": 2, "Strengthens": 3}
ROLE_COLORS = {
    "Core": "FFD54F",
    "Disease": "EF9A9A",
    "Treatment": "A5D6A7",
    "Strengthens": "90CAF9",
}

def add_cell(cell, text, bold=False, size=8, color=None, bg=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text))
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

doc = Document()

title = doc.add_heading("HeartOS — 50 Emotions & Their Related Attributes", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Complete attribute mapping: Core, Disease, Treatment & Strengthens relationships")
r.font.size = Pt(10)
r.italic = True

# Legend
legend = doc.add_paragraph()
legend.alignment = WD_ALIGN_PARAGRAPH.CENTER
legend_entries = [("Core", "FFD54F"), ("Disease", "EF9A9A"), ("Treatment", "A5D6A7"), ("Strengthens", "90CAF9")]
for label, color in legend_entries:
    run = legend.add_run(f"  {label}  ")
    run.font.size = Pt(8)
    run.bold = True
    from docx.oxml.ns import qn
    rPr = run._element.get_or_add_rPr()
    shd = rPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): color})
    rPr.append(shd)

doc.add_paragraph()

for idx, em in enumerate(emotions):
    eid = em["Emotion_ID"]
    h = doc.add_heading(f"{idx+1}. {em['Core_Emotion']} ({em.get('Arabic_Name','')})", level=1)

    em_links = links_by_emotion.get(eid, [])
    grouped = {}
    for link in em_links:
        role = link["Role"]
        grouped.setdefault(role, []).append(link)

    if not em_links:
        doc.add_paragraph("(No attributes linked)").paragraph_format.space_before = Pt(2)
        continue

    table = doc.add_table(rows=1, cols=6)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Role", "Attribute", "Arabic Name", "Urdu Name", "Nature", "Weight"]
    for ci, hdr in enumerate(headers):
        cell = table.cell(0, ci)
        add_cell(cell, hdr, bold=True, size=8, color=(255, 255, 255))
        from docx.oxml.ns import qn
        tcPr = cell._element.get_or_add_tcPr()
        shd = tcPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): '37474F'})
        tcPr.append(shd)

    for role in sorted(grouped.keys(), key=lambda r: ROLE_ORDER.get(r, 9)):
        for link in sorted(grouped[role], key=lambda x: x.get("Weight", 0), reverse=True):
            aid = link["Attribute_ID"]
            attr = attr_map.get(aid)
            if not attr:
                continue
            row = table.add_row()
            vals = [
                role,
                attr.get("Attribute", ""),
                attr.get("Arabic_Name", ""),
                attr.get("Urdu_Name", ""),
                attr.get("Nature", ""),
                str(link.get("Weight", "")),
            ]
            for ci, val in enumerate(vals):
                add_cell(row.cells[ci], val, size=8)
            # color the role cell
            add_cell(row.cells[0], role, bold=True, size=8)
            from docx.oxml.ns import qn
            tcPr = row.cells[0]._element.get_or_add_tcPr()
            shd = tcPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): ROLE_COLORS.get(role, "FFFFFF")})
            tcPr.append(shd)
            # color nature cell
            nature = attr.get("Nature", "")
            if nature == "Positive":
                tcPr2 = row.cells[4]._element.get_or_add_tcPr()
                shd2 = tcPr2.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'E8F5E9'})
                tcPr2.append(shd2)
            elif nature == "Negative":
                tcPr2 = row.cells[4]._element.get_or_add_tcPr()
                shd2 = tcPr2.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'FFEBEE'})
                tcPr2.append(shd2)

    for row in table.rows:
        row.cells[0].width = Inches(0.9)
        row.cells[1].width = Inches(1.5)
        row.cells[2].width = Inches(1.2)
        row.cells[3].width = Inches(1.0)
        row.cells[4].width = Inches(0.7)
        row.cells[5].width = Inches(0.5)

    doc.add_page_break()

output = os.path.join(BASE, "HeartOS_50_Emotions_Attributes_Reference.docx")
doc.save(output)
print(f"Saved: {output}")
print(f"Total emotions: {len(emotions)}")
total_links = sum(len(links_by_emotion.get(e["Emotion_ID"], [])) for e in emotions)
print(f"Total attribute links: {total_links}")
